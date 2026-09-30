import io

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from reportlab.lib import colors
from reportlab.lib.pagesizes import letter, landscape
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer


TITULOS = {
    "reservas": "Reporte de Reservas por Cancha",
    "ocupacion": "Reporte de Ocupación por Cancha",
    "ingresos": "Reporte de Ingresos por Cancha",
}


def _titulo(tipo):
    return TITULOS.get(tipo, f"Reporte: {tipo}")


def _formatear_valor(valor):
    if isinstance(valor, float):
        if abs(valor - round(valor)) < 1e-9:
            return str(int(valor))
        return f"{valor:.2f}"
    return str(valor)


def exportar_excel(tipo, filas, fecha_inicio, fecha_fin, nombre_empresa="Complejo Deportivo Canchas"):
    wb = Workbook()
    ws = wb.active
    ws.title = tipo[:31]

    ws["A1"] = nombre_empresa
    ws["A1"].font = Font(bold=True, size=14, color="1F2937")
    ws["A2"] = _titulo(tipo)
    ws["A2"].font = Font(bold=True, size=12)
    ws["A3"] = f"Desde {fecha_inicio} hasta {fecha_fin}"
    ws["A3"].font = Font(italic=True, size=10, color="6B7280")

    fila_actual = 5
    if not filas:
        ws.cell(row=fila_actual, column=1, value="Sin datos en el rango seleccionado.")
    else:
        encabezados = list(filas[0].keys())
        for col, enc in enumerate(encabezados, start=1):
            c = ws.cell(row=fila_actual, column=col, value=enc.replace("_", " ").title())
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = PatternFill("solid", fgColor="4F46E5")
            c.alignment = Alignment(horizontal="center", vertical="center")

        for fila in filas:
            fila_actual += 1
            for col, valor in enumerate(fila.values(), start=1):
                ws.cell(row=fila_actual, column=col, value=_formatear_valor(valor))

        for col in range(1, len(encabezados) + 1):
            ws.column_dimensions[ws.cell(row=5, column=col).column_letter].width = 22

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer


def exportar_pdf(tipo, filas, fecha_inicio, fecha_fin, nombre_empresa="Complejo Deportivo Canchas"):
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=landscape(letter),
        leftMargin=30,
        rightMargin=30,
        topMargin=30,
        bottomMargin=30,
    )
    estilos = getSampleStyleSheet()
    titulo = ParagraphStyle(
        "TituloR",
        parent=estilos["Title"],
        textColor=colors.HexColor("#111827"),
        fontSize=16,
    )
    sub = ParagraphStyle(
        "SubR",
        parent=estilos["Normal"],
        textColor=colors.HexColor("#6B7280"),
        fontSize=10,
    )

    elementos = [
        Paragraph(nombre_empresa, titulo),
        Paragraph(_titulo(tipo), estilos["Heading2"]),
        Paragraph(f"Desde {fecha_inicio} hasta {fecha_fin}", sub),
        Spacer(1, 12),
    ]

    if not filas:
        elementos.append(Paragraph("Sin datos en el rango seleccionado.", estilos["Normal"]))
    else:
        encabezados = list(filas[0].keys())
        data = [[e.replace("_", " ").title() for e in encabezados]] + [
            [_formatear_valor(v) for v in f.values()] for f in filas
        ]
        tabla = Table(data, repeatRows=1)
        tabla.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#4F46E5")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
            ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#CBD5E1")),
            ("FONTSIZE", (0, 0), (-1, -1), 8),
            ("ALIGN", (0, 0), (-1, 0), "CENTER"),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#EEF2FF")]),
        ]))
        elementos.append(tabla)

    doc.build(elementos)
    buffer.seek(0)
    return buffer